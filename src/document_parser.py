"""
Document Parser Module for Gemini TTS Studio
Extracts text from TXT, MD, PDF, DOCX, and SRT files with chapter splitting capabilities.
"""

import re
from pathlib import Path
from typing import List, Tuple, Optional


def extract_text_from_file(file_path: Path) -> str:
    """
    Extract clean plain text from a supported document file (.txt, .md, .pdf, .docx, .srt).
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"Datei nicht gefunden: {file_path}")

    ext = file_path.suffix.lower()

    if ext in [".txt", ".md"]:
        return _extract_from_text_file(file_path)
    elif ext == ".pdf":
        return _extract_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return _extract_from_docx(file_path)
    elif ext == ".srt":
        return _extract_from_srt(file_path)
    else:
        raise ValueError(f"Nicht unterstütztes Dateiformat: {ext} (Erlaubt: .txt, .md, .pdf, .docx, .srt)")


def _extract_from_text_file(file_path: Path) -> str:
    """Reads text file trying UTF-8, then fallback encodings."""
    encodings = ["utf-8", "utf-8-sig", "cp1252", "latin-1"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                return f.read().strip()
        except UnicodeDecodeError:
            continue
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read().strip()


def _extract_from_pdf(file_path: Path) -> str:
    """Extracts text from PDF pages using pypdf."""
    try:
        import pypdf
    except ImportError:
        raise ImportError("pypdf ist nicht installiert. Bitte 'pip install pypdf' ausführen.")

    reader = pypdf.PdfReader(str(file_path))
    extracted_pages = []
    
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        page_text = page_text.strip()
        if page_text:
            extracted_pages.append(page_text)

    full_text = "\n\n".join(extracted_pages)
    if not full_text.strip():
        raise ValueError(f"In der PDF '{file_path.name}' konnte kein lesbarer Text gefunden werden (evtl. gescanntes Bild).")
    
    return full_text


def _extract_from_docx(file_path: Path) -> str:
    """Extracts text from DOCX paragraphs and tables using python-docx."""
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx ist nicht installiert. Bitte 'pip install python-docx' ausführen.")

    doc = docx.Document(str(file_path))
    paragraphs = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            # Preserve heading structure if present
            if p.style and p.style.name.startswith("Heading"):
                paragraphs.append(f"\n# {text}\n")
            else:
                paragraphs.append(text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                paragraphs.append(" | ".join(row_texts))

    return "\n\n".join(paragraphs).strip()


def _extract_from_srt(file_path: Path) -> str:
    """Extracts speech text from SRT subtitle file, removing sequence numbers and timestamps."""
    raw_text = _extract_from_text_file(file_path)
    lines = raw_text.splitlines()
    clean_lines = []

    timestamp_pattern = re.compile(r"^\d{2}:\d{2}:\d{2}[,\.]\d{3}\s*-->\s*\d{2}:\d{2}:\d{2}[,\.]\d{3}")
    sequence_pattern = re.compile(r"^\d+$")

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if sequence_pattern.match(stripped):
            continue
        if timestamp_pattern.match(stripped):
            continue
        clean_lines.append(stripped)

    return " ".join(clean_lines).strip()


def split_into_chapters(text: str, default_title: str = "Kapitel") -> List[Tuple[str, str]]:
    """
    Splits a document text into chapters based on markdown headings (#, ##, ###)
    or common chapter headings ('Kapitel 1', 'Chapter 1', 'Teil 1').
    Returns a list of tuples: [(chapter_title, chapter_text), ...]
    """
    text = text.strip()
    if not text:
        return []

    # Regex patterns for chapter headers
    # 1. Markdown headings (# Kapitel ..., ## Einleitung, etc.)
    # 2. Plain text chapter headers (Kapitel 1: ..., Chapter 2 - ..., Teil 3 ...)
    heading_regex = re.compile(
        r"^(?:#{1,3}\s+(.+)|(?:Kapitel|Chapter|Teil|Abschnitt|Section)\s+([0-9IVXLCDM]+(?:\s*[:\-\.]\s*.+)?))\s*$",
        re.IGNORECASE | re.MULTILINE
    )

    matches = list(heading_regex.finditer(text))

    if not matches:
        # No clear chapter headers found -> return single document
        return [(default_title, text)]

    chapters: List[Tuple[str, str]] = []

    # If there is content before the first heading (e.g. Intro / Vorwort)
    first_match = matches[0]
    if first_match.start() > 0:
        pre_content = text[:first_match.start()].strip()
        if pre_content:
            chapters.append(("00_Einleitung", pre_content))

    for i, match in enumerate(matches):
        # Extract heading text
        heading_title = match.group(1) or match.group(2) or f"Kapitel_{i+1:02d}"
        heading_title = re.sub(r"[^\w\s\-_]", "", heading_title).strip()
        heading_title = re.sub(r"\s+", "_", heading_title)
        
        start_pos = match.end()
        end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        
        chapter_content = text[start_pos:end_pos].strip()
        if chapter_content:
            numbered_title = f"{len(chapters)+1:02d}_{heading_title[:40]}"
            chapters.append((numbered_title, chapter_content))

    return chapters if chapters else [(default_title, text)]
